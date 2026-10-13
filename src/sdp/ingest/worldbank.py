"""Monthly commodity prices and price indices from the World Bank Pink Sheet workbook.

    python -m sdp.ingest.worldbank             # pull when the table is a month old
    python -m sdp.ingest.worldbank --force     # pull now
    python -m sdp.ingest.worldbank --rebuild   # build again from the newest vendor pull

The workbook states the whole history from 1960 in nominal dollars and the Bank revises it, so
the table is current state. The commodity page links the workbook under a new address each month.
A series is a column of the sheet Monthly Prices (a price, with the unit of its header) or of
Monthly Indices (an index, 2010=100). Footnote stars and extra spaces are dropped from its
name. A cell with no value has no row. The month is the last day of the month.
vendor/ keeps the workbook of each pull.
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import logging
import os
import re
from pathlib import Path

import duckdb
import openpyxl

from sdp import dal
from sdp.config import settings
from sdp.ingest.common import AuditFailure, get_public, one, publish, rows_to_parquet
from sdp.ingest.rest import _temp_beside

log = logging.getLogger(__name__)

DATASET = dal.WORLDBANK_COMMODITIES

FILE = "CMO-Historical-Data-Monthly.xlsx"
LINK = re.compile(r'href="([^"]*' + re.escape(FILE) + r')"')
MONTH = re.compile(r"(\d{4})M(\d{2})")

PRICES = "Monthly Prices"
INDICES = "Monthly Indices"
INDEX_UNIT = "2010=100"

# Series name -> sheet. These are the columns kept; each must be in the workbook.
SERIES = {
    "Crude oil, average": PRICES,
    "Gold": PRICES,
    "Total Index": INDICES,
    "Energy": INDICES,
    "Non-energy": INDICES,
    "Agriculture": INDICES,
    "Metals & Minerals": INDICES,
}

COLUMNS = {"sheet": "VARCHAR", "series": "VARCHAR", "unit": "VARCHAR", "month": "DATE",
           "value": "DOUBLE", "vendor_pull_date": "DATE"}
MISSING = {"", "…", "..", "n/a"}

MAX_AGE_DAYS = 30
STALE_MONTHS = 3      # The Bank publishes in the first days of the next month; warn beyond this.

Rows = dict[str, tuple[str, list[tuple[str, float]]]]
"""series -> (unit, [(month as an ISO date, value)])."""


def _pull_dir(pull_date: dt.date) -> Path:
    return settings.vendor_dir / DATASET.name / f"{pull_date:%Y-%m-%d}"


def vendor_pulls() -> list[tuple[dt.date, Path]]:
    """Return the vendor pulls that hold the workbook, oldest first."""
    root = settings.vendor_dir / DATASET.name
    out = []
    for p in sorted(root.glob("????-??-??")) if root.exists() else []:
        try:
            d = dt.date.fromisoformat(p.name)
        except ValueError:
            continue
        if (p / FILE).exists():
            out.append((d, p / FILE))
    return out


def _download(url: str, attempts: int = 4) -> bytes:
    return get_public(url, attempts)


def latest_link(page: bytes) -> str:
    """Return the URL of the monthly workbook on the commodity page."""
    found = sorted(set(LINK.findall(page.decode("utf-8", "replace"))))
    if len(found) != 1:
        raise AuditFailure(f"The commodity page links {len(found)} monthly workbooks. "
                           f"The parser expects one.")
    return found[0]


def fetch(pull_date: dt.date) -> Path:
    """Download the workbook into the vendor directory of the pull date. Return its path."""
    url = latest_link(_download(settings.worldbank_page_url))
    data = _download(url)
    if not data.startswith(b"PK"):
        raise AuditFailure(f"{url} is not a workbook: {data[:40]!r}")
    d = _pull_dir(pull_date)
    d.mkdir(parents=True, exist_ok=True)
    tmp = _temp_beside(d / FILE)
    tmp.write_bytes(data)
    os.replace(tmp, d / FILE)
    log.info("Wrote %s bytes to %s", len(data), d / FILE)
    return d / FILE


def _name(cell) -> str:
    """Return a header without footnote stars and repeated spaces."""
    return re.sub(r"\s+", " ", str(cell).replace("*", "")).strip()


def _month(code: str) -> str:
    """Return the last day of a YYYYMmm month as an ISO date."""
    year, month = (int(x) for x in MONTH.fullmatch(code).groups())
    if not 1 <= month <= 12:
        raise AuditFailure(f"The workbook has the month {code}.")
    return f"{year:04d}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"


def _value(cell, series: str, code: str) -> float | None:
    if cell is None or (isinstance(cell, str) and cell.strip() in MISSING):
        return None
    try:
        return float(cell)
    except (TypeError, ValueError):
        raise AuditFailure(f"{series} has the value {cell!r} in {code}.") from None


def parse_sheet(rows: list[tuple], sheet: str) -> Rows:
    """Return the kept series of one sheet.

    The data rows start at the first row whose first cell is a month such as 1960M01. A price
    sheet has the names in the row two above it and the units in the row above. An index sheet
    has its names in a header of several rows, so a name is the last text of its column.
    """
    first = next((i for i, r in enumerate(rows)
                  if r and isinstance(r[0], str) and MONTH.fullmatch(r[0].strip())), None)
    if first is None or first < 2:
        raise AuditFailure(f"The sheet {sheet} has no monthly rows.")
    width = max(len(r) for r in rows[:first])
    names: dict[int, str] = {}
    units: dict[int, str] = {}
    if sheet == PRICES:
        for j in range(1, width):
            cell = rows[first - 2][j] if j < len(rows[first - 2]) else None
            if cell is not None and _name(cell):
                names[j] = _name(cell)
                unit = rows[first - 1][j] if j < len(rows[first - 1]) else None
                units[j] = str(unit).strip().strip("()").strip() if unit is not None else ""
    else:
        if not any(isinstance(r[0], str) and INDEX_UNIT in r[0] for r in rows[:first] if r):
            raise AuditFailure(f"The sheet {sheet} does not state its base, {INDEX_UNIT}.")
        for j in range(1, width):
            texts = [_name(r[j]) for r in rows[:first]
                     if j < len(r) and r[j] is not None and _name(r[j])]
            if texts:
                names[j], units[j] = texts[-1], INDEX_UNIT
    out: Rows = {}
    for series in (s for s, sh in SERIES.items() if sh == sheet):
        cols = [j for j, n in names.items() if n == series]
        if len(cols) != 1:
            raise AuditFailure(f"The sheet {sheet} has {len(cols)} columns named {series!r}.")
        j = cols[0]
        if not units[j]:
            raise AuditFailure(f"{series} has no unit.")
        values = []
        for r in rows[first:]:
            if not (r and isinstance(r[0], str) and MONTH.fullmatch(r[0].strip())):
                continue
            code = r[0].strip()
            v = _value(r[j] if j < len(r) else None, series, code)
            if v is not None:
                values.append((_month(code), v))
        out[series] = (units[j], values)
    return out


def read_workbook(path: Path) -> Rows:
    """Return the kept series of the workbook."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out: Rows = {}
        for sheet in dict.fromkeys(SERIES.values()):
            if sheet not in wb.sheetnames:
                raise AuditFailure(f"The workbook has no sheet {sheet!r}.")
            out.update(parse_sheet(list(wb[sheet].iter_rows(values_only=True)), sheet))
    finally:
        wb.close()
    return out


def build(src: Path, pull_date: dt.date) -> Path:
    """Build the table from one vendor workbook into _staging/. Return the staged file."""
    parsed = read_workbook(src)
    table = sorted((SERIES[s], s, unit, month, v, pull_date)
                   for s, (unit, values) in parsed.items() for month, v in values)
    staged = settings.staging_dir / f"{DATASET.name}-{os.getpid()}.parquet"
    rows_to_parquet(table, COLUMNS, staged)
    return staged


def _months_between(earlier: dt.date, later: dt.date) -> int:
    return (later.year - earlier.year) * 12 + later.month - earlier.month


def audit(staged: Path, pull_date: dt.date) -> int:
    """Raise AuditFailure on a table that would give wrong numbers; return the row count."""
    n, dups, nulls, nonpositive = one(f"""
        select count(*), count(*) - count(distinct (series, month)),
               count(*) filter (where series is null or month is null or value is null
                                or isnan(value)),
               count(*) filter (where value <= 0)
        from read_parquet('{staged}')""")
    fatal = []
    if dups:
        fatal.append(f"{dups} duplicate (series, month) rows")
    if nulls:
        fatal.append(f"{nulls} rows with a null key or value")
    if nonpositive:
        fatal.append(f"{nonpositive} prices or indices of 0 or less")
    last = dict(duckdb.execute(
        f"select series, max(month) from read_parquet('{staged}') group by 1").fetchall())
    absent = [s for s in SERIES if s not in last]
    if absent:
        fatal.append("no rows for " + ", ".join(absent))
    if DATASET.table_file.exists():
        (lost,) = one(f"""
            select count(*) from read_parquet('{DATASET.table_file}') p
            anti join read_parquet('{staged}') s on p.series = s.series and p.month = s.month""")
        if lost:
            fatal.append(f"{lost} months of the published table are not in the pull. "
                         f"The Bank does not remove history")
    if fatal:
        raise AuditFailure(f"{DATASET.name}: " + ". ".join(fatal) + ".")
    for series, month in sorted(last.items()):
        if _months_between(month, pull_date) > STALE_MONTHS:
            log.warning("%s ends in %s, more than %s months before the pull.", series, month,
                        STALE_MONTHS)
    return n


def published_pull() -> dt.date | None:
    """Return the vendor pull date of the published table, or None."""
    if not DATASET.table_file.exists():
        return None
    (d,) = one(f"select max(vendor_pull_date) from read_parquet('{DATASET.table_file}')")
    return d


def ingest(pull_date: dt.date | None = None, *, force: bool = False,
           rebuild: bool = False) -> Path:
    """Pull the workbook and publish the table.

    Without force, do nothing while the published table is younger than MAX_AGE_DAYS.
    rebuild reads the newest vendor pull without downloading.
    """
    pull_date = pull_date or dt.datetime.now(dt.UTC).date()
    last = published_pull()
    if not (force or rebuild) and last and (pull_date - last).days < MAX_AGE_DAYS:
        log.info("%s is from the pull of %s. The next pull is due %s days later.",
                 DATASET.name, last, MAX_AGE_DAYS)
        return DATASET.table_file
    if rebuild:
        pulls = vendor_pulls()
        if not pulls:
            raise FileNotFoundError(f"{DATASET.name} has no vendor pull to rebuild from.")
        pull_date, src = pulls[-1]
    else:
        src = fetch(pull_date)
    staged = build(src, pull_date)
    try:
        n = audit(staged, pull_date)
    except BaseException:
        staged.unlink(missing_ok=True)
        raise
    path = publish(staged, DATASET.table_file)
    log.info("Published %s: %s rows from the pull of %s.", path, n, pull_date)
    return path


def main(argv: list[str] | None = None) -> None:
    p = argparse.ArgumentParser(prog="python -m sdp.ingest.worldbank",
                                description="Pull the World Bank monthly commodity prices.")
    mode = p.add_mutually_exclusive_group()
    mode.add_argument("--force", action="store_true", help="Pull now, whatever the age.")
    mode.add_argument("--rebuild", action="store_true",
                      help="Build again from the newest vendor pull. Do not download.")
    args = p.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    print(ingest(force=args.force, rebuild=args.rebuild))


if __name__ == "__main__":
    main()
