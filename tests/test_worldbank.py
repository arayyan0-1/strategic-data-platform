"""Tests for the World Bank commodity ingest against a synthetic Pink Sheet workbook."""
import datetime as dt
import io
import logging

import openpyxl
import pytest

from sdp import dal
from sdp.config import settings
from sdp.ingest import worldbank as wbk
from sdp.ingest.common import AuditFailure

PULL = dt.date(2022, 7, 4)
N = 30   # Months from 2020-01 to 2022-06.
URL = "https://docs.example.org/en/doc/abc123-0050012022/related/" + wbk.FILE
PAGE = (f'<a href="{URL}">Monthly prices</a><a href="https://docs.example.org/x/CMO-Historical-'
        f'Data-Annual.xlsx">Annual prices</a>').encode()


def _code(i: int) -> str:
    return f"{2020 + i // 12}M{i % 12 + 1:02d}"


def workbook(n: int = N, *, bad: str | None = None, base: float = 0.0) -> bytes:
    """A Pink Sheet in miniature: prices with a names row and a units row, and indices under a
    header of several rows with footnote stars and a double space."""
    wb = openpyxl.Workbook()
    wb.active.title = "Mismatch Details"
    prices = wb.create_sheet("Monthly Prices")
    prices.append(["World Bank Commodity Price Data (The Pink Sheet)"])
    prices.append(["monthly prices in nominal US dollars, 1960 to present"])
    prices.append([None])
    prices.append(["Updated on July 02, 2022"])
    gold = "Gold bars" if bad == "missing_series" else "Gold"
    prices.append([None, "Crude oil, average", "Crude oil, Brent", gold, "Beef **"])
    prices.append([None, "($/bbl)", "($/bbl)", "($/troy oz)", "($/kg)"])
    for i in range(n):
        row = [_code(i), 50 + i + base, 60 + i, 1500 + 10 * i + base, 5.0]
        if i == 0:
            row[3] = "…"
        if bad == "text_value" and i == 5:
            row[3] = "abc"
        if bad == "zero" and i == 5:
            row[3] = 0
        if bad == "bad_month" and i == 7:
            row[0] = "2020M13"
        prices.append(row)
    if bad == "duplicate":
        prices.append([_code(n - 1), 1, 2, 3, 4])
    if bad != "no_sheet":
        indices = wb.create_sheet("Monthly Indices")
        base_line = ("monthly indices based on nominal US dollars, 2010=100, 1960 to present"
                     if bad != "no_base" else "monthly indices, 1960 to present")
        indices.append([None])
        indices.append([base_line])
        indices.append([None])
        indices.append(["Updated on July 02, 2022"])
        indices.append([None])
        indices.append([None, "Total Index", "Energy", "Non-energy **"])
        indices.append([None, None, None, None, "Agriculture **", None, "Metals  & Minerals"])
        indices.append([None, None, None, None, None, "Beverages"])
        indices.append([None, None, None, " ", " "])
        for i in range(n):
            indices.append([_code(i), 100 + i + base, 90 + i, 80 + i, 70 + i, 50, 60 + i])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def site(tmp_data_root, monkeypatch):
    """Serve the page and a workbook. Return the served files and the download log."""
    served: dict[str, bytes] = {settings.worldbank_page_url: PAGE, URL: workbook()}
    calls: list[str] = []

    def download(url, attempts=4):
        calls.append(url)
        return served[url]

    monkeypatch.setattr(wbk, "_download", download)
    return served, calls


def _rows(series: str):
    return dal.worldbank_commodities().filter(f"series = '{series}'").order("month").fetchall()


def test_a_pull_publishes_the_kept_series_with_units_on_the_last_day_of_the_month(site):
    _, calls = site
    assert wbk.ingest(PULL) == dal.WORLDBANK_COMMODITIES.table_file
    assert calls == [settings.worldbank_page_url, URL]
    names = dal.worldbank_commodities().aggregate("series, count(*)").order("series").fetchall()
    assert [r[0] for r in names] == sorted(wbk.SERIES), "the stars and the extra space are gone"
    sheet, series, unit, month, value, pull = _rows("Crude oil, average")[1]
    assert (sheet, unit, month, value, pull) == ("Monthly Prices", "$/bbl", dt.date(2020, 2, 29),
                                                 51.0, PULL), "February 2020 has 29 days"
    assert _rows("Total Index")[0][2] == "2010=100"
    assert _rows("Metals & Minerals")[0][4] == 60.0


def test_a_cell_without_a_value_has_no_row(site):
    wbk.ingest(PULL)
    gold = _rows("Gold")
    assert len(gold) == N - 1
    assert gold[0][3] == dt.date(2020, 2, 29)
    assert gold[0][2] == "$/troy oz"
    assert len(_rows("Energy")) == N


def test_a_series_that_is_not_kept_has_no_row(site):
    wbk.ingest(PULL)
    assert dal.worldbank_commodities().filter(
        "series like 'Beef%' or series like 'Crude oil, Brent%' or series = 'Beverages'"
    ).count("*").fetchone() == (0,)


def test_a_fresh_table_is_not_pulled_again(site):
    _, calls = site
    wbk.ingest(PULL)
    wbk.ingest(PULL + dt.timedelta(days=wbk.MAX_AGE_DAYS - 1))
    assert len(calls) == 2
    wbk.ingest(PULL + dt.timedelta(days=wbk.MAX_AGE_DAYS))
    assert len(calls) == 4


def test_a_rebuild_reads_the_vendor_workbook_and_does_not_download(site):
    _, calls = site
    wbk.ingest(PULL)
    dal.WORLDBANK_COMMODITIES.table_file.unlink()
    wbk.ingest(rebuild=True)
    assert len(calls) == 2
    assert len(_rows("Gold")) == N - 1
    assert _rows("Gold")[0][5] == PULL


@pytest.mark.parametrize("bad", ["duplicate", "missing_series", "text_value", "zero",
                                 "bad_month", "no_base", "no_sheet"])
def test_a_bad_workbook_is_not_published(site, bad):
    served, _ = site
    served[URL] = workbook(bad=bad)
    with pytest.raises(AuditFailure):
        wbk.ingest(PULL)
    assert not dal.WORLDBANK_COMMODITIES.table_file.exists()


def test_a_duplicate_month_is_named_in_the_error(site):
    served, _ = site
    served[URL] = workbook(bad="duplicate")
    with pytest.raises(AuditFailure, match="duplicate"):
        wbk.ingest(PULL)


def test_a_page_with_no_workbook_link_or_two_is_refused(site):
    served, _ = site
    served[settings.worldbank_page_url] = b"<html>no links</html>"
    with pytest.raises(AuditFailure, match="0 monthly workbooks"):
        wbk.ingest(PULL)
    served[settings.worldbank_page_url] = PAGE + PAGE.replace(b"abc123", b"def456")
    with pytest.raises(AuditFailure, match="2 monthly workbooks"):
        wbk.ingest(PULL)


def test_a_download_that_is_not_a_workbook_is_refused(site):
    served, _ = site
    served[URL] = b"<html>Access denied</html>"
    with pytest.raises(AuditFailure, match="not a workbook"):
        wbk.ingest(PULL)
    assert not dal.WORLDBANK_COMMODITIES.table_file.exists()


def test_a_pull_that_loses_a_month_keeps_the_old_table(site):
    served, _ = site
    wbk.ingest(PULL)
    before = dal.WORLDBANK_COMMODITIES.table_file.read_bytes()
    served[URL] = workbook(n=N - 3)
    with pytest.raises(AuditFailure, match="does not remove history"):
        wbk.ingest(PULL, force=True)
    assert dal.WORLDBANK_COMMODITIES.table_file.read_bytes() == before


def test_a_restated_value_and_a_new_month_are_published(site):
    served, _ = site
    wbk.ingest(PULL)
    served[URL] = workbook(n=N + 2, base=1.0)
    wbk.ingest(PULL, force=True)
    total = _rows("Total Index")
    assert len(total) == N + 2
    assert total[0][4] == 101.0, "the Bank restated January 2020"
    assert total[-1][3] == dt.date(2022, 8, 31)


def test_a_series_that_ends_long_before_the_pull_warns(site, caplog):
    with caplog.at_level(logging.WARNING, logger=wbk.log.name):
        wbk.ingest(dt.date(2022, 9, 20))
    assert "ends in 2022-06-30" not in caplog.text, "three months is not stale"
    with caplog.at_level(logging.WARNING, logger=wbk.log.name):
        wbk.ingest(dt.date(2022, 10, 20), force=True)
    assert "Gold ends in 2022-06-30" in caplog.text
